# certificates.py - Chapter 19 mini project
"""Mail-merge: one Word certificate per row of an Excel sheet, plus a combined PDF and a log."""
import csv
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

BASE = Path(__file__).resolve().parent
ROSTER = BASE / "roster.xlsx"
OUT = BASE / "certificates"


def make_sample_roster() -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Course", "Hours", "Score"])
    for row in [("Ana Cruz", "Python Foundations", 24, 96), ("Ben Santos", "Python Foundations", 24, 88),
                ("Cara Reyes", "Data Analysis with pandas", 16, 91), ("Dan Lim", "Python Foundations", 24, 72)]:
        ws.append(row)
    wb.save(ROSTER)


def read_roster(path: Path) -> list[dict]:
    ws = load_workbook(path, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() for h in next(rows)]
    return [dict(zip(header, r)) for r in rows if r[0]]


def word_certificate(p: dict, folder: Path) -> Path:
    doc = Document()
    def line(text, size, bold=False, color=(0, 0, 0), space=12):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(space)
        run = para.add_run(text)
        run.font.size, run.bold = Pt(size), bold
        run.font.color.rgb = RGBColor(*color)
    line("CERTIFICATE OF COMPLETION", 26, True, (31, 78, 121), 30)
    line("This certifies that", 14)
    line(p["name"], 30, True, space=18)
    line(f"has completed {p['course']} ({p['hours']} hours)", 14)
    line(f"with a final score of {p['score']}%" + (" - with honors" if p["score"] >= 90 else ""), 14, space=40)
    line(f"Issued {date.today():%d %B %Y}", 11, color=(110, 110, 110))
    path = folder / f"{p['name'].replace(' ', '_')}.docx"
    doc.save(path)
    return path


def pdf_certificates(people: list[dict], path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=landscape(A4))
    w, h = landscape(A4)
    for p in people:
        c.setStrokeColorRGB(0.12, 0.31, 0.47)
        c.setLineWidth(4)
        c.rect(30, 30, w - 60, h - 60)
        c.setFont("Helvetica-Bold", 30)
        c.drawCentredString(w / 2, h - 130, "CERTIFICATE OF COMPLETION")
        c.setFont("Helvetica", 16)
        c.drawCentredString(w / 2, h - 190, "This certifies that")
        c.setFont("Helvetica-Bold", 34)
        c.drawCentredString(w / 2, h - 250, p["name"])
        c.setFont("Helvetica", 16)
        c.drawCentredString(w / 2, h - 300, f"has completed {p['course']} ({p['hours']} hours)")
        c.drawCentredString(w / 2, h - 330, f"final score {p['score']}%")
        c.setFont("Helvetica", 11)
        c.drawCentredString(w / 2, 70, f"Issued {date.today():%d %B %Y}")
        c.showPage()                       # one page per person
    c.save()


if __name__ == "__main__":
    if not ROSTER.exists():
        make_sample_roster()
    people = [p for p in read_roster(ROSTER) if p["score"] >= 75]      # passing only
    OUT.mkdir(exist_ok=True)
    with open(OUT / "issued.csv", "w", newline="", encoding="utf-8") as f:
        log = csv.writer(f)
        log.writerow(["name", "course", "file"])
        for p in people:
            log.writerow([p["name"], p["course"], word_certificate(p, OUT).name])
    pdf_certificates(people, OUT / "all_certificates.pdf")
    print(f"{len(people)} certificates written to {OUT.name}/ (Word files + one PDF)")
