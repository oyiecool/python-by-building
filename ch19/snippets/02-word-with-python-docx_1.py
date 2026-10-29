from docx import Document
from docx.shared import Cm

doc = Document()
doc.add_heading("Incident report: web01 outage", level=1)
p = doc.add_paragraph("On 28 October the web server was unreachable for ")
p.add_run("42 minutes").bold = True
p.add_run(" after the disk filled up.")
doc.add_heading("Timeline", level=2)
for t, event in [("02:10", "Disk reached 100%"), ("02:14", "Monitoring alert sent"), ("02:52", "Logs rotated, service restored")]:
    doc.add_paragraph(f"{t}  {event}", style="List Bullet")
table = doc.add_table(rows=1, cols=3, style="Light Grid Accent 1")
for cell, text in zip(table.rows[0].cells, ["Action", "Owner", "Due"]):
    cell.text = text
row = table.add_row().cells
row[0].text, row[1].text, row[2].text = "Add disk alert at 85%", "Ronald", "2026-11-05"
doc.add_picture("metrics_report.png", width=Cm(15))   # any PNG, e.g. from Capstone 3
doc.save("incident_report.docx")

# reading an existing document
for para in Document("incident_report.docx").paragraphs:
    print(para.style.name, "|", para.text)
