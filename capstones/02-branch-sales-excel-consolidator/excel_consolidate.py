# excel_consolidate.py - Capstone 2
"""Merge branch sales workbooks into one formatted report with formulas and a chart."""
import random
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
INBOX = BASE / "branch_reports"
OUT = BASE / "consolidated_sales.xlsx"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")


def make_sample_inputs() -> None:
    """Create one workbook per branch, the way branches would email them in."""
    INBOX.mkdir(exist_ok=True)
    random.seed(3)
    products = ["Laptop", "Phone", "Tablet", "Monitor", "Headset"]
    for branch in ["Manila", "Cebu", "Davao", "Baguio"]:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Product", *MONTHS])
        scale = {"Manila": 1.6, "Cebu": 1.1, "Davao": 0.9, "Baguio": 0.6}[branch]
        for p in products:
            ws.append([p, *[round(random.uniform(80, 220) * scale) * 100 for _ in MONTHS]])
        wb.save(INBOX / f"{branch}_sales.xlsx")


def read_branch(path: Path) -> list[list]:
    ws = load_workbook(path, data_only=True)["Sales"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return [list(r) for r in rows if r[0]]


def style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center")


def autosize(ws) -> None:
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 3, 40)


def build() -> None:
    files = sorted(INBOX.glob("*_sales.xlsx"))
    if not files:
        raise SystemExit(f"No branch files in {INBOX}")
    wb = Workbook()

    # 1. One combined data sheet: Branch | Product | months... | Total (formula)
    data = wb.active
    data.title = "All branches"
    data.append(["Branch", "Product", *MONTHS, "Total"])
    for f in files:
        branch = f.stem.replace("_sales", "")
        for row in read_branch(f):
            r = data.max_row + 1
            data.append([branch, *row])
            first, last = get_column_letter(3), get_column_letter(2 + len(MONTHS))
            data.cell(row=r, column=3 + len(MONTHS), value=f"=SUM({first}{r}:{last}{r})")
    style_header(data, 3 + len(MONTHS))
    for row in data.iter_rows(min_row=2, min_col=3):
        for c in row:
            c.number_format = "#,##0"
            c.border = Border(bottom=THIN)
    data.freeze_panes = "C2"
    data.auto_filter.ref = data.dimensions
    autosize(data)

    # 2. Summary sheet driven by SUMIF formulas, so it updates if data changes
    summ = wb.create_sheet("Summary", 0)
    summ.append(["Branch", "Total sales", "Share", "Target", "vs target"])
    branches = [f.stem.replace("_sales", "") for f in files]
    total_col = get_column_letter(3 + len(MONTHS))
    n = data.max_row
    for i, b in enumerate(branches, start=2):
        summ.append([
            b,
            f"=SUMIF('All branches'!A2:A{n},A{i},'All branches'!{total_col}2:{total_col}{n})",
            f"=B{i}/B{len(branches) + 2}",
            500000 if b != "Manila" else 900000,
            f"=B{i}/D{i}-1",
        ])
    t = len(branches) + 2
    summ.append(["TOTAL", f"=SUM(B2:B{t - 1})", f"=SUM(C2:C{t - 1})", f"=SUM(D2:D{t - 1})", f"=B{t}/D{t}-1"])
    style_header(summ, 5)
    for c in summ[t]:
        c.font = Font(bold=True)
    for r in range(2, t + 1):
        summ[f"B{r}"].number_format = summ[f"D{r}"].number_format = "#,##0"
        summ[f"C{r}"].number_format = summ[f"E{r}"].number_format = "0.0%"
    red = PatternFill("solid", fgColor="F8CBAD")
    green = PatternFill("solid", fgColor="C6EFCE")
    summ.conditional_formatting.add(f"E2:E{t}", CellIsRule(operator="lessThan", formula=["0"], fill=red))
    summ.conditional_formatting.add(f"E2:E{t}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=green))
    autosize(summ)

    # 3. A native Excel bar chart of sales by branch
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Total sales by branch"
    chart.y_axis.title = "Sales"
    chart.add_data(Reference(summ, min_col=2, min_row=1, max_row=t - 1), titles_from_data=True)
    chart.set_categories(Reference(summ, min_col=1, min_row=2, max_row=t - 1))
    chart.height, chart.width = 8, 16
    summ.add_chart(chart, "G2")

    wb.save(OUT)
    print(f"Merged {len(files)} branch files, {n - 1} product rows -> {OUT.name}")


if __name__ == "__main__":
    if not INBOX.exists():
        make_sample_inputs()
        print(f"Created sample branch files in {INBOX.name}/")
    build()
